"""Export test cases to various formats (CSV, Excel, DataFrame)."""

import io
from typing import Any, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class TestCase:
    """Represents a single test case."""

    def __init__(
        self,
        tc_id: str,
        title: str,
        case_type: str,
        priority: str,
        preconditions: str,
        steps: List[str],
        expected_result: str,
        ac_reference: str,
    ):
        self.tc_id = tc_id
        self.title = title
        self.case_type = case_type
        self.priority = priority
        self.preconditions = preconditions
        self.steps = steps
        self.expected_result = expected_result
        self.ac_reference = ac_reference


def to_dataframe(test_cases: List[TestCase]) -> pd.DataFrame:
    """Convert test cases to a pandas DataFrame for display."""
    data = []
    for tc in test_cases:
        steps_text = "\n".join(f"{i+1}. {step}" for i, step in enumerate(tc.steps))
        data.append({
            "TC ID": tc.tc_id,
            "Title": tc.title,
            "Type": tc.case_type,
            "Priority": tc.priority,
            "Preconditions": tc.preconditions,
            "Steps": steps_text,
            "Expected Result": tc.expected_result,
            "AC Reference": tc.ac_reference,
        })
    return pd.DataFrame(data)


def to_csv_bytes(test_cases: List[TestCase]) -> bytes:
    """Convert test cases to CSV bytes."""
    df = to_dataframe(test_cases)
    csv_buffer = io.StringIO()
    df.to_csv(csv_buffer, index=False)
    return csv_buffer.getvalue().encode("utf-8")


def to_excel_bytes(test_cases: List[TestCase]) -> bytes:
    """Convert test cases to Excel bytes with formatting."""
    df = to_dataframe(test_cases)

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Test Cases", index=False)

        # Format the workbook
        workbook = writer.book
        worksheet = writer.sheets["Test Cases"]

        # Header formatting
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Column widths
        worksheet.column_dimensions["A"].width = 12  # TC ID
        worksheet.column_dimensions["B"].width = 25  # Title
        worksheet.column_dimensions["C"].width = 12  # Type
        worksheet.column_dimensions["D"].width = 12  # Priority
        worksheet.column_dimensions["E"].width = 20  # Preconditions
        worksheet.column_dimensions["F"].width = 30  # Steps
        worksheet.column_dimensions["G"].width = 25  # Expected Result
        worksheet.column_dimensions["H"].width = 15  # AC Reference

        # Data cell formatting
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Freeze header row
        worksheet.freeze_panes = "A2"

    return excel_buffer.getvalue()
